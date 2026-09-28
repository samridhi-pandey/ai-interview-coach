"""Module 1: CV text extraction + transparent, rule-based ATS scoring (+ AI written feedback when available)."""
import io
import re

import ai_service
from models import CVNarrative
from prompts import CV_SYSTEM

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_CV_CHARS = 20000  # far above a normal CV; keeps prompts bounded

# Category weights (points) - sum to 100.
WEIGHTS = {"structure": 20, "formatting": 20, "content": 25, "keywords": 20, "completeness": 15}


class CVError(Exception):
    pass


# ---------------------------------------------------------------- extraction
def extract_text(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if len(data) > MAX_FILE_BYTES:
        raise CVError("File is too large (max 5 MB).")
    try:
        if name.endswith(".pdf"):
            import pymupdf  # PyMuPDF

            with pymupdf.open(stream=data, filetype="pdf") as doc:
                text = "\n".join(page.get_text() for page in doc)
        elif name.endswith(".docx"):
            import docx

            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.extend(cell.text for cell in row.cells)
            text = "\n".join(parts)
        else:
            raise CVError("Unsupported file type. Please upload a PDF or DOCX file.")
    except CVError:
        raise
    except Exception:
        raise CVError("Could not read this file. Make sure it is a valid, non-password-protected PDF or DOCX.")

    text = re.sub(r"[ \t]+", " ", text).strip()
    if len(text) < 80:
        raise CVError(
            "Very little text could be extracted. If your CV is a scanned image, "
            "please upload a text-based PDF or a DOCX file."
        )
    return text[:MAX_CV_CHARS]


# ---------------------------------------------------------------- detection
SECTION_PATTERNS = {
    "Summary/Objective": r"summary|objective|profile|about me|career objective",
    "Education": r"education|academic|qualification",
    "Skills": r"skills?|core competenc\w+|technologies|tech stack|technical proficiency",
    "Projects": r"projects?",
    "Experience": r"experience|internships?|employment|work history",
    "Certifications": r"certifications?|certificates?|courses|training",
    "Achievements": r"achievements?|awards?|honou?rs|accomplishments|extracurricular|activities|positions? of responsibility",
}

SKILLS = [
    # languages
    "Python", "Java", "JavaScript", "TypeScript", "C++", "C#", "C", "Golang", "Rust", "Kotlin", "Swift", "PHP", "Ruby",
    "MATLAB", "Scala", "Dart", "Bash", "SQL", "HTML", "CSS",
    # web / backend
    "React", "Angular", "Vue", "Node.js", "Express.js", "Django", "Flask", "FastAPI", "Spring Boot", ".NET",
    "Next.js", "Tailwind", "Bootstrap", "REST", "GraphQL", "Flutter", "Android", "jQuery",
    # data / ML
    "Machine Learning", "Deep Learning", "NLP", "Computer Vision", "Data Science", "Data Analysis", "TensorFlow",
    "PyTorch", "Keras", "scikit-learn", "Pandas", "NumPy", "Matplotlib", "OpenCV", "Power BI", "Tableau", "Excel",
    "Generative AI", "LLM", "Hugging Face", "LangChain", "Statistics",
    # databases
    "MySQL", "PostgreSQL", "MongoDB", "SQLite", "Oracle", "Redis", "Firebase", "DBMS", "NoSQL",
    # tools / cloud / devops
    "Git", "GitHub", "Docker", "Kubernetes", "AWS", "Azure", "GCP", "Linux", "Jenkins", "CI/CD", "Postman", "Jira",
    "VS Code", "Figma", "Selenium", "Hadoop", "Spark", "Kafka",
    # CS fundamentals
    "Data Structures", "Algorithms", "DSA", "OOP", "Object-Oriented Programming", "Operating Systems",
    "Computer Networks", "System Design", "Agile", "Scrum", "Cybersecurity", "IoT", "Embedded Systems", "Arduino",
    "Raspberry Pi", "Problem Solving",
]

# Short/ambiguous names that are also normal English words - match exact case only.
CASE_SENSITIVE = {"Excel", "Oracle", "Spark", "Swift", "Scala", "Statistics"}

ACTION_VERBS = {
    "developed", "built", "designed", "implemented", "created", "led", "managed", "improved", "optimized", "optimised",
    "reduced", "increased", "achieved", "analyzed", "analysed", "automated", "deployed", "integrated", "collaborated",
    "delivered", "engineered", "launched", "trained", "tested", "coordinated", "organized", "organised", "presented",
    "designed", "architected", "resolved", "streamlined", "established", "conducted", "won", "secured", "researched",
    "applied", "programmed", "maintained", "migrated", "enhanced", "mentored", "published", "constructed", "generated",
}

BULLET_RE = re.compile(r"^\s*[•●▪■◦\-\*–·►➢✓]\s*\S")
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,5}\)?[\s-]?)?\d{3,5}[\s-]?\d{4,5}")
LINK_RE = re.compile(r"linkedin\.com|github\.com|gitlab\.com|portfolio|leetcode\.com|kaggle\.com", re.I)
DEGREE_RE = re.compile(r"b\.?\s?tech|b\.?\s?e\b|bachelor|m\.?\s?tech|master|b\.?\s?sc|m\.?\s?sc|bca|mca|b\.?\s?com|mba|diploma|cgpa|gpa", re.I)
METRIC_RE = re.compile(
    r"\d+(?:\.\d+)?\s?%"
    r"|\$\s?\d[\d,.]*"
    r"|\b\d[\d,]*\+?\s?(?:users|customers|students|records|requests|apis?|models?|members|participants|ms|seconds|hours|"
    r"teams?|datasets?|samples|images|endpoints|features|lines|stars|downloads|x)\b"
    r"|\b(?:increased|reduced|improved|decreased|boosted|cut|saved|achieved|ranked|secured)\b[^.\n]{0,40}\d",
    re.I,
)
FIRST_PERSON_RE = re.compile(r"\b(?:I|my|me|myself)\b")


def detect_sections(text: str) -> list[str]:
    found = []
    for line in text.splitlines():
        line = line.strip()
        if not line or len(line) > 45 or line.endswith("."):
            continue
        norm = re.sub(r"[^a-z/ ]", " ", line.lower())
        norm = re.sub(r"\s+", " ", norm).strip()
        if not norm or len(norm.split()) > 5:
            continue
        for section, pat in SECTION_PATTERNS.items():
            if section not in found and re.search(rf"\b(?:{pat})\b", norm):
                found.append(section)
    order = list(SECTION_PATTERNS)
    return sorted(found, key=order.index)


def detect_skills(text: str) -> list[str]:
    found = []
    for skill in SKILLS:
        # custom boundaries so "C++", "C#", ".NET", "Node.js" work
        pat = rf"(?<![A-Za-z0-9+#]){re.escape(skill)}(?![A-Za-z0-9+#])"
        if re.search(pat, text, re.I if len(skill) > 4 and skill not in CASE_SENSITIVE else 0):
            found.append(skill)
    return found


# ---------------------------------------------------------------- scoring
def _scaled(value: float, full_at: float, max_pts: int) -> int:
    return round(min(value, full_at) / full_at * max_pts)


def score_resume(text: str) -> dict:
    """Deterministic ATS scoring. Every point is traceable to something found (or missing) in the CV text."""
    lines = [l for l in text.splitlines() if l.strip()]
    words = re.findall(r"\b\w+\b", text)
    sections = detect_sections(text)
    skills = detect_skills(text)
    has = lambda s: s in sections

    has_email = bool(EMAIL_RE.search(text))
    phone_ok = any(len(re.sub(r"\D", "", m.group())) >= 10 for m in PHONE_RE.finditer(text))
    has_links = bool(LINK_RE.search(text))
    bullets = [l for l in lines if BULLET_RE.match(l)]
    words_lower = {w.lower() for w in words}
    verbs = sorted(ACTION_VERBS & words_lower)
    metrics = METRIC_RE.findall(text)
    long_bullets = [b for b in bullets if len(b.split()) >= 8]
    fp_count = len(FIRST_PERSON_RE.findall(text))

    # Structure (20)
    structure = 0
    structure += 5 if has("Education") else 0
    structure += 5 if has("Skills") else 0
    structure += 6 if has("Projects") and has("Experience") else 4 if (has("Projects") or has("Experience")) else 0
    structure += min(4, 2 * sum(has(s) for s in ("Summary/Objective", "Certifications", "Achievements")))

    # Formatting (20)
    n = len(words)
    length_pts = 5 if 250 <= n <= 900 else 3 if 150 <= n < 250 or 900 < n <= 1200 else 1
    formatting = (4 if has_email else 0) + (4 if phone_ok else 0) + (3 if has_links else 0) + length_pts
    formatting += _scaled(len(bullets), 6, 4)  # uses bullet points

    # Content (25)
    content = _scaled(len(verbs), 8, 8) + _scaled(len(metrics), 5, 9) + _scaled(len(long_bullets), 5, 4)
    content += 4 if fp_count <= 2 else 2 if fp_count <= 6 else 0

    # Keywords (20)
    keywords = _scaled(len(skills), 12, 20)

    # Completeness (15)
    completeness = (
        (2 if has_email else 0)
        + (2 if phone_ok else 0)
        + (2 if has("Education") else 0)
        + (1 if DEGREE_RE.search(text) else 0)
        + (3 if has("Skills") else 0)
        + (3 if has("Projects") or has("Experience") else 0)
        + (2 if has("Summary/Objective") or has_links else 0)
    )

    cats = {
        "structure": structure,
        "formatting": formatting,
        "content": content,
        "keywords": keywords,
        "completeness": completeness,
    }
    for k in cats:
        cats[k] = max(0, min(cats[k], WEIGHTS[k]))

    facts = dict(
        word_count=n, has_email=has_email, has_phone=phone_ok, has_links=has_links, bullet_count=len(bullets),
        action_verbs=verbs, metric_count=len(metrics), first_person=fp_count,
    )
    return {
        "ats_score": sum(cats.values()),
        "category_scores": cats,
        "category_max": WEIGHTS,
        "category_percent": {k: round(cats[k] / WEIGHTS[k] * 100) for k in cats},
        "detected_sections": sections,
        "detected_skills": skills,
        "facts": facts,
    }


def rule_based_narrative(r: dict) -> CVNarrative:
    f, secs, skills = r["facts"], r["detected_sections"], r["detected_skills"]
    strengths, gaps, recs = [], [], []

    if len(skills) >= 8:
        strengths.append(f"Good technical keyword coverage - {len(skills)} relevant skills detected.")
    if "Projects" in secs:
        strengths.append("Includes a Projects section, which is valuable for internship and placement CVs.")
    if "Experience" in secs:
        strengths.append("Includes work/internship experience.")
    if f["metric_count"] >= 3:
        strengths.append(f"Uses measurable results ({f['metric_count']} quantified statements found).")
    if len(f["action_verbs"]) >= 6:
        strengths.append("Uses strong action verbs to describe work.")
    if f["has_email"] and f["has_phone"]:
        strengths.append("Clear contact details (email and phone) are present.")
    if not strengths:
        strengths.append("The CV has readable, extractable text that ATS systems can parse.")

    for sec in ("Education", "Skills", "Projects"):
        if sec not in secs:
            gaps.append(f"No clearly labelled '{sec}' section was detected.")
            recs.append(f"Add a clearly titled '{sec}' section so ATS software can find it.")
    if f["metric_count"] < 3:
        gaps.append("Few or no measurable achievements.")
        recs.append("Add numbers to your project/experience bullets (e.g. accuracy %, users, speed-up, dataset size).")
    if len(skills) < 8:
        gaps.append(f"Limited keyword coverage - only {len(skills)} recognised skills.")
        recs.append("List the relevant tools, languages and CS fundamentals you genuinely know, using standard names.")
    if len(f["action_verbs"]) < 6:
        gaps.append("Few action verbs in the descriptions.")
        recs.append("Start bullets with action verbs such as Developed, Implemented, Designed or Optimized.")
    if f["bullet_count"] < 5:
        gaps.append("Little use of bullet points.")
        recs.append("Use short bullet points instead of long paragraphs for projects and experience.")
    if not f["has_links"]:
        recs.append("Consider adding your GitHub or LinkedIn link.")
    if not (f["has_email"] and f["has_phone"]):
        gaps.append("Contact details (email/phone) are missing or unclear.")
        recs.append("Put your email and phone number at the top of the CV.")
    if f["word_count"] > 900:
        recs.append("Consider shortening the CV - one page is ideal for students.")
    if not gaps:
        gaps.append("No major structural gaps detected by the automated check.")
    if not recs:
        recs.append("Tailor the skills and project descriptions to each role you apply for.")
    return CVNarrative(strengths=strengths[:5], areas_to_improve=gaps[:5], recommendations=recs[:5])


def analyze(text: str) -> dict:
    result = score_resume(text)
    mode, note = "demo", None
    narrative = rule_based_narrative(result)
    if ai_service.AI_ENABLED:
        user = (
            f"ATS check results (rule-based): score {result['ats_score']}/100; categories {result['category_scores']} "
            f"out of {WEIGHTS}; sections found: {result['detected_sections']}; skills found: {result['detected_skills']}; "
            f"facts: {result['facts']}\n\n<candidate_cv>\n{text}\n</candidate_cv>"
        )
        try:
            narrative = ai_service.generate_json(CV_SYSTEM, user, CVNarrative, max_tokens=3000)
            mode = "ai"
        except ai_service.AIError as e:
            note = f"AI feedback unavailable ({e}). Showing rule-based feedback instead."
    result.pop("facts")
    result.update(
        strengths=narrative.strengths,
        areas_to_improve=narrative.areas_to_improve,
        recommendations=narrative.recommendations,
        mode=mode,
        note=note,
    )
    return result
